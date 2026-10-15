"""Сервис для парсинга банковских выписок (CSV, XLSX)."""
from __future__ import annotations

import csv
import io
import logging
from datetime import datetime, date
from decimal import Decimal, InvalidOperation
from typing import List, Dict, Any, Optional

import openpyxl

logger = logging.getLogger(__name__)

# Возможные названия колонок (нижний регистр, без пробелов)
DATE_COLUMNS = {"дата", "date", "oper_date", "operationdate", "датаоперации", "датасписания"}
AMOUNT_COLUMNS = {"сумма", "amount", "sum", "суммаоперации", "sumoperation", "суммасписания"}
DESC_COLUMNS = {"описание", "description", "назначение", "назначениеплатежа", "details", "комментарий", "comment", "получатель", "плательщик"}
CATEGORY_COLUMNS = {"категория", "category", "mcc", " mcc "}

def _normalize_header(h: str) -> str:
    """Приводит заголовок к нижнему регистру и удаляет пробелы."""
    return h.lower().replace(" ", "").replace("_", "")

def _parse_date(val: str) -> Optional[date]:
    """Пытается распарсить дату из различных форматов."""
    if not val:
        return None
    val = val.strip()
    formats = [
        "%d.%m.%Y", "%d.%m.%y", "%Y-%m-%d", "%d/%m/%Y", "%Y/%m/%d",
        "%d.%m.%Y %H:%M:%S", "%Y-%m-%d %H:%M:%S"
    ]
    for fmt in formats:
        try:
            return datetime.strptime(val, fmt).date()
        except ValueError:
            continue
    return None

def _parse_amount(val: str) -> Optional[Decimal]:
    """Парсит сумму, заменяя запятые на точки и обрабатывая минусы."""
    if not val:
        return None
    val = str(val).strip().replace(" ", "").replace("\xa0", "").replace(",", ".")
    # Убираем символы валют
    val = "".join(c for c in val if c.isdigit() or c in ".-")
    try:
        return Decimal(val)
    except InvalidOperation:
        return None

def _find_column_index(headers: List[str], target_names: set) -> Optional[int]:
    """Находит индекс колонки по возможным названиям."""
    norm_headers = [_normalize_header(h) for h in headers]
    for target in target_names:
        norm_target = _normalize_header(target)
        if norm_target in norm_headers:
            return norm_headers.index(norm_target)
    return None

def parse_csv_file(file_content: bytes) -> List[Dict[str, Any]]:
    """Парсит CSV файл (пробует UTF-8, затем CP1251)."""
    rows = []
    for encoding in ["utf-8-sig", "utf-8", "cp1251"]:
        try:
            text = file_content.decode(encoding)
            break
        except UnicodeDecodeError:
            continue
    else:
        raise ValueError("Не удалось определить кодировку CSV файла. Попробуйте сохранить в UTF-8.")

    # Пробуем разные разделители
    for delimiter in [",", ";", "\t"]:
        try:
            reader = csv.DictReader(io.StringIO(text), delimiter=delimiter)
            headers = reader.fieldnames or []
            if not headers:
                continue
            
            idx_date = _find_column_index(headers, DATE_COLUMNS)
            idx_amount = _find_column_index(headers, AMOUNT_COLUMNS)
            idx_desc = _find_column_index(headers, DESC_COLUMNS)
            idx_cat = _find_column_index(headers, CATEGORY_COLUMNS)

            if idx_date is None or idx_amount is None:
                continue # Не те колонки, пробуем другой разделитель

            for row in reader:
                header_list = list(row.keys())
                date_val = row[header_list[idx_date]] if idx_date < len(header_list) else ""
                amount_val = row[header_list[idx_amount]] if idx_amount < len(header_list) else ""
                desc_val = row[header_list[idx_desc]] if idx_desc is not None and idx_desc < len(header_list) else ""
                cat_val = row[header_list[idx_cat]] if idx_cat is not None and idx_cat < len(header_list) else ""

                parsed_date = _parse_date(date_val)
                parsed_amount = _parse_amount(amount_val)

                if parsed_date and parsed_amount is not None:
                    rows.append({
                        "date": parsed_date.isoformat(),
                        "amount": float(parsed_amount),
                        "description": desc_val.strip() if desc_val else "",
                        "category_hint": cat_val.strip() if cat_val else "",
                    })
            
            if rows:
                return rows
        except Exception:
            continue
            
    raise ValueError("Не удалось распарсить CSV. Проверьте формат и наличие колонок 'Дата' и 'Сумма'.")

def parse_xlsx_file(file_content: bytes) -> List[Dict[str, Any]]:
    """Парсит XLSX файл."""
    rows = []
    try:
        wb = openpyxl.load_workbook(io.BytesIO(file_content), read_only=True, data_only=True)
        sheet = wb.active
        if not sheet:
            raise ValueError("Файл Excel пуст или не содержит активных листов.")
        
        # Читаем первые 2 строки для определения заголовков
        header_row = next(sheet.iter_rows(min_row=1, max_row=1, values_only=True))
        if not header_row or not any(header_row):
            # Иногда заголовки во второй строке
            header_row = next(sheet.iter_rows(min_row=2, max_row=2, values_only=True))
            
        headers = [str(h) if h is not None else "" for h in header_row]
        
        idx_date = _find_column_index(headers, DATE_COLUMNS)
        idx_amount = _find_column_index(headers, AMOUNT_COLUMNS)
        idx_desc = _find_column_index(headers, DESC_COLUMNS)
        idx_cat = _find_column_index(headers, CATEGORY_COLUMNS)

        if idx_date is None or idx_amount is None:
            raise ValueError(f"Не найдены обязательные колонки 'Дата' или 'Сумма'. Найдены: {', '.join(headers)}")

        for row in sheet.iter_rows(min_row=2, values_only=True):
            if not any(row): # Пустая строка
                continue
                
            row_list = list(row)
            date_val = str(row_list[idx_date]) if idx_date < len(row_list) else ""
            amount_val = str(row_list[idx_amount]) if idx_amount < len(row_list) else ""
            desc_val = str(row_list[idx_desc]) if idx_desc is not None and idx_desc < len(row_list) else ""
            cat_val = str(row_list[idx_cat]) if idx_cat is not None and idx_cat < len(row_list) else ""

            parsed_date = _parse_date(date_val)
            parsed_amount = _parse_amount(amount_val)

            if parsed_date and parsed_amount is not None:
                rows.append({
                    "date": parsed_date.isoformat(),
                    "amount": float(parsed_amount),
                    "description": desc_val.strip() if desc_val and desc_val != "None" else "",
                    "category_hint": cat_val.strip() if cat_val and cat_val != "None" else "",
                })
        
        return rows
    except Exception as e:
        raise ValueError(f"Ошибка парсинга Excel: {str(e)}")

def parse_statement(file_content: bytes, filename: str) -> List[Dict[str, Any]]:
    """Определяет тип файла и вызывает соответствующий парсер."""
    if filename.lower().endswith(".csv"):
        return parse_csv_file(file_content)
    elif filename.lower().endswith((".xlsx", ".xls")):
        return parse_xlsx_file(file_content)
    else:
        raise ValueError("Неподдерживаемый формат файла. Используйте CSV или XLSX.")
